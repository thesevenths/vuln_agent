#!/usr/bin/env python3
"""生成 GuardFox Agent 宏观架构 PowerPoint（无代码细节）。"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUTPUT_CN = ROOT / "docs" / "GuardFox-Agent三阶段架构.pptx"
OUTPUT_EN = ROOT / "docs" / "guardfox-agent-3stage-architecture.pptx"

C_TITLE = RGBColor(0x1E, 0x29, 0x3B)
C_ACCENT = RGBColor(0x1A, 0x56, 0x9B)
C_STAGE1 = RGBColor(0xDA, 0xE8, 0xFC)
C_STAGE2 = RGBColor(0xD5, 0xE8, 0xD4)
C_STAGE3 = RGBColor(0xFF, 0xE6, 0xCC)
C_ORCH = RGBColor(0xE1, 0xD5, 0xE7)
C_PLAN = RGBColor(0xF5, 0xF5, 0xF5)
C_BORDER1 = RGBColor(0x6C, 0x8E, 0xBF)
C_BORDER2 = RGBColor(0x82, 0xB3, 0x66)
C_BORDER3 = RGBColor(0xD6, 0xB6, 0x56)
C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def _set_slide_title(slide, title: str) -> None:
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.1))
    bar.fill.solid()
    bar.fill.fore_color.rgb = C_ACCENT
    bar.line.fill.background()
    box = slide.shapes.add_textbox(Inches(0.4), Inches(0.15), Inches(12.5), Inches(0.6))
    p = box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = C_TITLE


def _box(slide, left, top, width, height, text, fill, border, font_size=9, bold=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = border
    shape.line.width = Pt(1)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.alignment = PP_ALIGN.CENTER
    return shape


def _arrow_down(slide, x, y1, y2):
    slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x, y1, x, y2).line.color.rgb = C_ACCENT


def slide_architecture(prs):
    """幻灯片 1：宏观架构设计（对齐参考图）。"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_title(slide, "一、架构设计 — AI 白盒漏洞挖掘三阶段方案")

    _box(
        slide, Inches(0.4), Inches(0.85), Inches(5.8), Inches(0.55),
        "编排器\n阶段调度 · 状态持久化 · 异常处理/熔断",
        C_ORCH, C_ACCENT, 10, True,
    )
    _box(
        slide, Inches(6.4), Inches(0.85), Inches(6.5), Inches(0.55),
        "数据存储\n源码 · 知识库 · 策略 · 审计日志 · 报告/PoC",
        C_ORCH, C_ACCENT, 10, True,
    )
    _arrow_down(slide, Inches(6.65), Inches(1.45), Inches(1.55))

    _box(
        slide, Inches(0.4), Inches(1.55), Inches(12.5), Inches(0.32),
        "阶段一：多源并行侦查  →  统一候选 UFF",
        C_STAGE1, C_BORDER1, 11, True,
    )
    ch_w = Inches(2.95)
    y1 = Inches(1.95)
    h1 = Inches(1.35)
    _box(slide, Inches(0.45), y1, ch_w, h1,
         "CPG 污点流追踪\nSQLi/命令/XSS/SSRF\n反序列化/路径穿越\n产出：污点路径", C_WHITE, C_BORDER1, 9)
    _box(slide, Inches(3.55), y1, ch_w, h1,
         "结构/内存缺陷检测\n溢出/UAF/Double-Free\n嵌入式 CAN/MMIO\n产出：触发路径", C_WHITE, C_BORDER1, 9)
    _box(slide, Inches(6.65), y1, ch_w, h1,
         "LLM 语义逻辑审计\n鉴权/越权/IDOR\n会话/业务流\n产出：逻辑检查项", C_WHITE, C_BORDER1, 9)
    _box(slide, Inches(9.75), y1, ch_w, h1,
         "规则 SAST 接入\nSemgrep/SARIF\nCI快扫候选\n不直接定性", C_PLAN, C_BORDER1, 9)

    _arrow_down(slide, Inches(6.65), Inches(3.35), Inches(3.48))

    _box(
        slide, Inches(0.4), Inches(3.48), Inches(12.5), Inches(0.32),
        "阶段二：深度审计 / 可利用性分析",
        C_STAGE2, C_BORDER2, 11, True,
    )
    y2 = Inches(3.9)
    w2 = Inches(3.9)
    h2 = Inches(1.1)
    _box(slide, Inches(0.45), y2, w2, h2,
         "安全缺陷验证 Agent\n可达·路径·Sink·消毒", C_WHITE, C_BORDER2, 9)
    _box(slide, Inches(4.55), y2, w2, h2,
         "结构缺陷验证 Agent\n触发条件·防护缺失", C_WHITE, C_BORDER2, 9)
    _box(slide, Inches(8.65), y2, w2, h2,
         "逻辑缺陷验证 Agent\n鉴权链·业务检查单", C_WHITE, C_BORDER2, 9)
    _box(
        slide, Inches(0.45), Inches(5.15), Inches(12.45), Inches(0.42),
        "裁决：EXPLOITABLE | UNCERTAIN | FALSE_POSITIVE | FAILED  ·  证据：L1代码流 L2配置 L3动态",
        C_WHITE, C_BORDER2, 9,
    )

    _arrow_down(slide, Inches(6.65), Inches(5.62), Inches(5.75))

    _box(
        slide, Inches(0.4), Inches(5.75), Inches(12.5), Inches(0.32),
        "阶段三：利用验证引擎（双模式）",
        C_STAGE3, C_BORDER3, 11, True,
    )
    _box(slide, Inches(0.45), Inches(6.15), Inches(5.9), Inches(0.72),
         "模式 A · 服务交互\nWeb/API · 部署探测 · 会话攻击链", C_WHITE, C_BORDER3, 9)
    _box(slide, Inches(6.55), Inches(6.15), Inches(5.9), Inches(0.72),
         "模式 B · Harness\n库/内核/CLI · 触发输入 · ASan/崩溃", C_WHITE, C_BORDER3, 9)

    _box(
        slide, Inches(0.45), Inches(7.0), Inches(12.45), Inches(0.38),
        "最终交付：分级报告 · 攻击链可视化 · 可选 PoC · 修复建议",
        C_ACCENT, C_ACCENT, 10, True,
    )
    slide.shapes[-1].text_frame.paragraphs[0].font.color.rgb = C_WHITE


def slide_implementation(prs):
    """幻灯片 2：如何实现（分期 + 模块能力）。"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_title(slide, "二、如何实现 — 分期建设与功能模块")

    phases = [
        ("一期（当前）", "阶段一 A/B/C 并行侦查\n阶段二 三类验证 + 四态裁决\n编排器 + 审计 + 分级报告", C_STAGE2, C_BORDER2),
        ("二期", "SAST/SARIF 接入\nUFF 统一入库与去重\nCI 快扫 + Agent 验证闭环", C_STAGE1, C_BORDER1),
        ("三期", "双模式 PoC 引擎\nL3 证据回写\n攻击链可视化 / CVSS", C_STAGE3, C_BORDER3),
    ]
    x = Inches(0.45)
    for title, body, fill, border in phases:
        _box(slide, x, Inches(1.0), Inches(3.95), Inches(1.35), f"{title}\n{body}", fill, border, 10, True)
        x += Inches(4.1)

    _box(
        slide, Inches(0.45), Inches(2.55), Inches(12.4), Inches(0.65),
        "阶段衔接：任务接入 → 编排器选场景 Skill → 阶段一并行 → UFF 入库 → 阶段二分流验证 → [可选] 阶段三证明 → 报告",
        C_WHITE, C_ACCENT, 10,
    )

    modules = [
        ("工作流 / Skill 注册", C_ORCH, C_ACCENT),
        ("CPG 构建 / 查询包", C_STAGE1, C_BORDER1),
        ("智能验证 / 反证", C_STAGE2, C_BORDER2),
        ("知识库 / Playbook", C_STAGE1, C_BORDER1),
        ("证据补全 / 攻击面", C_STAGE2, C_BORDER2),
        ("动态验证引擎", C_STAGE3, C_BORDER3),
    ]
    x = Inches(0.45)
    y = Inches(3.45)
    mw = Inches(1.95)
    for name, fill, border in modules:
        _box(slide, x, y, mw, Inches(0.85), name, fill, border, 9, True)
        x += mw + Inches(0.12)

    _box(
        slide, Inches(0.45), Inches(4.55), Inches(12.4), Inches(0.9),
        "与 GuardFox 产品能力映射\n"
        "· C/C++、Java 编译深扫 → 阶段一 A/B + 阶段二\n"
        "· 非编译问答/预存储 → 数据层索引 + 可升级深扫\n"
        "· SCA/Scanner/SBOM → 通道 D + 供应链可达 Skill",
        C_PLAN, C_BORDER1, 10,
    )


def slide_advantages(prs):
    """幻灯片 3：核心优势。"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_title(slide, "三、核心优势")

    items = [
        ("候选→验证→裁决", "非告警堆砌\n四态可决策结论"),
        ("反证降噪", "消毒/配置核查\n过滤 SAST 误报"),
        ("证据三层", "L1/L2/L3\n不越级宣称可利用"),
        ("内存/嵌入式", "结构通道差异化\nCAN·MMIO·OTA"),
        ("逻辑补盲", "鉴权/业务专通道\nSAST 天然弱项"),
        ("全链路审计", "可回放可评测\n策略包版本锁定"),
    ]
    x0, y0 = Inches(0.45), Inches(1.0)
    bw, bh, gap = Inches(3.95), Inches(1.3), Inches(0.15)
    fills = [C_STAGE1, C_STAGE2, C_STAGE3, C_STAGE1, C_STAGE2, C_STAGE3]
    borders = [C_BORDER1, C_BORDER2, C_BORDER3, C_BORDER1, C_BORDER2, C_BORDER3]
    for i, (title, body) in enumerate(items):
        _box(slide, x0 + (i % 3) * (bw + gap), y0 + (i // 3) * (bh + gap), bw, bh,
             f"{title}\n{body}", fills[i], borders[i], 10, True)

    _box(
        slide, Inches(0.45), Inches(4.1), Inches(12.4), Inches(1.0),
        "相对传统 SAST：广撒网在 SAST，深验证在 GuardFox\n"
        "相对参考三阶段：三类验证 Agent + 统一裁决，非百个空 Agent；壁垒在阶段二证据体系\n"
        "定位：多源候选 → Agent 证伪分级 → 可选动态证明",
        C_ACCENT, C_ACCENT, 11, True,
    )
    slide.shapes[-1].text_frame.paragraphs[0].font.color.rgb = C_WHITE


def slide_sast(prs):
    """幻灯片 4：SAST 能发现的漏洞。"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_title(slide, "四、SAST 能发现的漏洞（通道 D）")

    _box(
        slide, Inches(0.45), Inches(1.0), Inches(5.9), Inches(2.2),
        "规则 SAST 擅长（模式/语法可匹配）\n\n"
        "· 注入：SQLi、命令、LDAP、XPath、日志注入\n"
        "· 输出：XSS 模式、响应头注入\n"
        "· 文件：路径穿越、不安全文件操作\n"
        "· 密码学：弱算法、硬编码密钥\n"
        "· API：危险函数、不安全反序列化模式\n"
        "· 配置：CORS、调试开关、敏感信息泄露",
        C_STAGE1, C_BORDER1, 10,
    )
    _box(
        slide, Inches(6.55), Inches(1.0), Inches(5.9), Inches(2.2),
        "SAST 边界（需 CPG 或逻辑通道补位）\n\n"
        "· 污点是否真正到达 Sink → CPG 污点流\n"
        "· 鉴权/过滤器是否生效 → 逻辑审计 + L2\n"
        "· 业务状态机/支付逻辑 → 逻辑专审\n"
        "· 内存二次释放/竞态 → 结构/内存通道\n"
        "· CVE 版本受影响 ≠ 可达 → 供应链可达",
        C_PLAN, C_BORDER1, 10,
    )
    _box(
        slide, Inches(0.45), Inches(3.45), Inches(12.4), Inches(0.75),
        "产品策略：SAST 负责广撒网（UFF 候选）→ 阶段二 Agent 负责证伪与分级 → 不把 SAST 告警当最终结论",
        C_WHITE, C_BORDER2, 11,
    )


def slide_business_logic(prs):
    """幻灯片 5：业务逻辑漏洞。"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_title(slide, "五、业务逻辑漏洞（通道 C + 逻辑验证 Agent）")

    headers = ["类型", "分析思路", "证据要求"]
    rows = [
        ["认证绕过", "攻击面测绘 + 鉴权链（过滤器/注解/白名单）", "L2 配置 + L3 复核"],
        ["水平/垂直越权", "对象 ID 与 Session 主体校验", "L2 + 业务检查单"],
        ["IDOR", "接口参数跨用户访问资源", "L1 线索 + L2 + L3"],
        ["会话管理缺陷", "Cookie/Token 策略、固定会话", "L2"],
        ["支付/价格篡改", "金额、状态字段来源与校验", "L3 必验"],
        ["工作流跳过", "状态机转移条件", "检查单 + L3"],
        ["并发/重复提交", "幂等、锁、双花窗口", "L3"],
    ]
    col_w = [Inches(2.5), Inches(5.5), Inches(2.3)]
    x0, y0 = Inches(0.45), Inches(1.0)
    row_h = Inches(0.48)
    x = x0
    for i, h in enumerate(headers):
        _box(slide, x, y0, col_w[i], row_h, h, C_ACCENT, C_ACCENT, 10, True)
        slide.shapes[-1].text_frame.paragraphs[0].font.color.rgb = C_WHITE
        x += col_w[i] + Inches(0.05)
    for ri, row in enumerate(rows):
        y = y0 + row_h + Inches(0.06) + ri * (row_h + Inches(0.04))
        x = x0
        fill = C_WHITE if ri % 2 == 0 else C_PLAN
        for ci, cell in enumerate(row):
            _box(slide, x, y, col_w[ci], row_h, cell, fill, C_BORDER2, 9)
            x += col_w[ci] + Inches(0.05)

    _box(
        slide, Inches(0.45), Inches(5.35), Inches(12.4), Inches(0.85),
        "专审流水线（逻辑专审 Skill）：攻击面测绘 → 鉴权配置分析 → 逻辑检查单 → 逻辑验证 Agent → 裁决\n"
        "原则：无 L2 不得 EXPLOITABLE；业务类无 L3 最高 UNCERTAIN；输出须结构化检查单（角色/对象/状态/幂等）",
        C_STAGE2, C_BORDER2, 10,
    )


def slide_tech_and_skills(prs):
    """幻灯片 6：技术方案与 Skill 设计。"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_title(slide, "六、技术方案与 Skill 设计")

    layers = [
        ("L2 领域知识 Skill", "鉴权手册 · 路由映射 · SQL/内存审计 · 业务逻辑检查单", C_STAGE2, C_BORDER2),
        ("L1 场景流水线 Skill", "源码深扫 · Web全链路 · 逻辑专审 · SAST验证 · 供应链可达 · 版本对齐 · 动态PoC", C_STAGE1, C_BORDER1),
        ("L0 原子能力 Skill", "CPG分析 · 配置读取 · 攻击面测绘 · 知识检索 · 版本识别 · SCA", C_WHITE, C_BORDER1),
        ("L3 裁决与证据", "硬回溯 · 反证复核 · L1/L2/L3 分级 · 终态与 PoC 策略", C_STAGE3, C_BORDER3),
    ]
    y = Inches(1.0)
    for title, body, fill, border in layers:
        _box(slide, Inches(0.45), y, Inches(12.4), Inches(0.95), f"{title}\n{body}", fill, border, 10)
        y += Inches(1.05)

    _box(
        slide, Inches(0.45), Inches(5.35), Inches(12.4), Inches(1.0),
        "技术方案要点\n"
        "侦查：四通道并行 · 查询/策略版本化 · 嵌入式分 profile\n"
        "归一：UFF 入库 · 多源去重\n"
        "验证：三类 Agent 分流 · 反证优先 · 证据不足不升格\n"
        "证明：双模式选型 · PoC 与报告绑定\n\n"
        "Skill 公式：触发条件 + 固定阶段 + 绑定能力 + 交付物 + 完成标准",
        C_ORCH, C_ACCENT, 10,
    )


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide_architecture(prs)
    slide_implementation(prs)
    slide_advantages(prs)
    slide_sast(prs)
    slide_business_logic(prs)
    slide_tech_and_skills(prs)
    OUTPUT_CN.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUTPUT_CN))
    prs.save(str(OUTPUT_EN))
    return OUTPUT_CN, OUTPUT_EN


if __name__ == "__main__":
    paths = build()
    for path in paths:
        print(f"Generated: {path}")
